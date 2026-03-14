from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QLabel, QTabWidget, QTableWidgetItem,
    QHBoxLayout, QPushButton, QFileDialog, QFrame, QMessageBox,
    QComboBox, QGridLayout
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QFont, QColor
from PyQt6.QtWidgets import QHeaderView

from ui.common.cards import CardFrame
from ui.common.tables import StyledTableWidget
from ui.common.styled_dialogs import info, warning, error
from ui.theme import heading_font, Theme, body_font
from api.global_client import api_client


class ReportsScreen(QWidget):
    def __init__(self, parent=None):
        super().__init__()
        self.parent_window = parent
        self.attendance_data = []
        self.semester_map = {}
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 24, 24, 24)
        root.setSpacing(20)

        title = QLabel("📄 Reports")
        title.setFont(heading_font(24, bold=True))
        title.setStyleSheet(f"color: {Theme.text_primary}; margin-bottom: 4px;")
        root.addWidget(title)

        actions_card = CardFrame(padding=16)
        actions = QHBoxLayout()
        actions.setSpacing(12)

        refresh_btn = QPushButton("🔄 Refresh Data")
        refresh_btn.setStyleSheet(self._btn(Theme.secondary, "#16A085"))
        refresh_btn.clicked.connect(self.load_reports)

        actions.addWidget(refresh_btn)
        actions.addStretch()
        actions_card.layout.addLayout(actions)
        root.addWidget(actions_card)

        self.tabs = QTabWidget()
        self.tabs.setStyleSheet(f"""
            QTabWidget::pane {{
                border: 1px solid {Theme.border};
                border-radius: 8px;
                background: white;
            }}
            QTabBar::tab {{
                padding: 10px 20px;
                font-size: 13px;
                font-weight: 600;
                color: {Theme.text_muted};
                border-bottom: 2px solid transparent;
            }}
            QTabBar::tab:selected {{
                color: {Theme.primary};
                border-bottom: 2px solid {Theme.primary};
            }}
        """)

        self.tabs.addTab(self._create_attendance_tab(), "📊 Attendance")
        self.tabs.addTab(self._create_viva_tab(), "🎤 Viva Marks")
        self.tabs.addTab(self._create_submissions_tab(), "📝 Reports")

        root.addWidget(self.tabs)
        root.addStretch(1)

    def _btn(self, bg, hover):
        return f"""
            QPushButton {{
                background: {bg};
                color: white;
                border: none;
                border-radius: 6px;
                padding: 10px 20px;
                font-weight: 600;
                font-size: 13px;
            }}
            QPushButton:hover {{
                background: {hover};
            }}
            QPushButton:pressed {{
                background: {hover}cc;
            }}
        """

    # =========================================================
    # TAB 1 — Attendance
    # =========================================================

    def _create_attendance_tab(self):
        widget = CardFrame(padding=20)

        filters_layout = QGridLayout()
        filters_layout.setHorizontalSpacing(12)
        filters_layout.setVerticalSpacing(8)

        sem_lbl = QLabel("Semester")
        sem_lbl.setFont(body_font(12, QFont.Weight.DemiBold))
        filters_layout.addWidget(sem_lbl, 0, 0)
        self.semester_combo = QComboBox()
        self.semester_combo.addItem("Select semester...", None)
        self.semester_combo.currentIndexChanged.connect(self.on_semester_changed)
        filters_layout.addWidget(self.semester_combo, 1, 0)

        subject_lbl = QLabel("Subject")
        subject_lbl.setFont(body_font(12, QFont.Weight.DemiBold))
        filters_layout.addWidget(subject_lbl, 0, 1)
        self.subject_combo = QComboBox()
        self.subject_combo.addItem("Select subject...", "")
        filters_layout.addWidget(self.subject_combo, 1, 1)

        self.attendance_load_btn = QPushButton("Show Attendance Report")
        self.attendance_load_btn.clicked.connect(self.load_attendance_report)
        filters_layout.addWidget(self.attendance_load_btn, 1, 2)

        widget.layout.addLayout(filters_layout)

        self.attendance_summary_label = QLabel("Select semester and subject to view attendance percentage.")
        self.attendance_summary_label.setFont(body_font(12))
        self.attendance_summary_label.setStyleSheet(
            f"color: {Theme.text_secondary}; padding: 8px 12px; background: {Theme.background}; border-radius: 8px;"
        )
        widget.layout.addWidget(self.attendance_summary_label)

        self.attendance_table = StyledTableWidget(0, 5)
        self.attendance_table.setHorizontalHeaderLabels([
            "Student Name",
            "Student ID",
            "Present / Total",
            "Attendance %",
            "Batch",
        ])
        widget.layout.addWidget(self.attendance_table)
        return widget
    
    def load_reports(self):
        """Load attendance filter data and refresh report if filters are selected."""
        timetable_result = api_client.get_faculty_weekly_timetable()
        if not timetable_result["success"]:
            QMessageBox.warning(self, "Error", f"Failed to load timetable filters:\n{timetable_result['error']}")
            return

        payload = timetable_result.get("data") if isinstance(timetable_result.get("data"), dict) else {}
        semesters = payload.get("semesters", [])
        slots = payload.get("slots", [])

        self.semester_map = {}
        for sem in semesters:
            sem_id = sem.get("id")
            if not sem_id:
                continue
            self.semester_map[sem_id] = {
                "id": sem_id,
                "name": sem.get("name", "Semester"),
                "number": sem.get("number", 0),
                "batches": sem.get("batches", []),
                "subjects": set(),
            }

        for slot in slots:
            sem = slot.get("semester", {})
            sem_id = sem.get("id")
            subject = (slot.get("subject_name") or "").strip()
            if sem_id in self.semester_map and subject:
                self.semester_map[sem_id]["subjects"].add(subject)

        previous_sem_id = self.semester_combo.currentData() if hasattr(self, "semester_combo") else None
        previous_subject = self.subject_combo.currentData() if hasattr(self, "subject_combo") else ""
        self.semester_combo.blockSignals(True)
        self.semester_combo.clear()
        self.semester_combo.addItem("Select semester...", None)
        for sem in sorted(self.semester_map.values(), key=lambda s: (s["number"], s["name"])):
            self.semester_combo.addItem(sem["name"], sem["id"])
        self.semester_combo.blockSignals(False)

        if previous_sem_id in self.semester_map:
            idx = self.semester_combo.findData(previous_sem_id)
            if idx >= 0:
                self.semester_combo.setCurrentIndex(idx)
                self.on_semester_changed(idx, preferred_subject=previous_subject)
                if self.subject_combo.currentData():
                    self.load_attendance_report()
                if hasattr(self.parent_window, 'current_batch_id') and self.parent_window.current_batch_id is not None:
                    self.load_submission_students()
                return

        self.on_semester_changed(self.semester_combo.currentIndex())
        if hasattr(self.parent_window, 'current_batch_id') and self.parent_window.current_batch_id is not None:
            self.load_submission_students()

    def on_semester_changed(self, _index: int, preferred_subject: str = ""):
        sem_id = self.semester_combo.currentData()
        self.subject_combo.clear()
        self.subject_combo.addItem("Select subject...", "")
        self.attendance_table.setRowCount(0)

        if not sem_id or sem_id not in self.semester_map:
            self.attendance_summary_label.setText("Select semester and subject to view attendance percentage.")
            return

        subjects = sorted(self.semester_map[sem_id]["subjects"])
        for subject in subjects:
            self.subject_combo.addItem(subject, subject)

        if preferred_subject:
            subject_index = self.subject_combo.findData(preferred_subject)
            if subject_index >= 0:
                self.subject_combo.setCurrentIndex(subject_index)

        if subjects:
            self.attendance_summary_label.setText("Select a subject and click 'Show Attendance Report'.")
        else:
            self.attendance_summary_label.setText("No subjects found for selected semester.")

    def load_attendance_report(self):
        sem_id = self.semester_combo.currentData()
        subject_name = (self.subject_combo.currentData() or "").strip()

        if not sem_id:
            QMessageBox.information(self, "Semester Required", "Please select a semester.")
            return
        if not subject_name:
            QMessageBox.information(self, "Subject Required", "Please select a subject.")
            return

        sem_data = self.semester_map.get(sem_id)
        if not sem_data:
            QMessageBox.warning(self, "Error", "Semester metadata not available.")
            return

        attendance_result = api_client.get_semester_attendance(semester_id=int(sem_id), subject_name=subject_name)
        if not attendance_result["success"]:
            QMessageBox.warning(self, "Error", f"Failed to load attendance report:\n{attendance_result['error']}")
            return

        records = attendance_result.get("data") or []
        session_ids = {r.get("session") for r in records if r.get("session")}
        total_sessions = len(session_ids)

        students_by_id = {}
        for batch in sem_data.get("batches", []):
            batch_id = batch.get("id")
            if not batch_id:
                continue
            students_result = api_client.get_students(batch_id=batch_id)
            if not students_result["success"]:
                continue
            for st in students_result.get("data") or []:
                sid = st.get("id")
                if not sid:
                    continue
                students_by_id[sid] = {
                    "id": sid,
                    "name": st.get("name", "Unknown"),
                    "student_id": st.get("student_id", "N/A"),
                    "batch_name": st.get("batch_name", "-"),
                }

        present_counts = {}
        for rec in records:
            sid = rec.get("student")
            if not sid:
                continue
            if rec.get("status") == "present":
                present_counts[sid] = present_counts.get(sid, 0) + 1

        sorted_students = sorted(students_by_id.values(), key=lambda s: s["name"].lower())
        self.attendance_table.setRowCount(len(sorted_students))

        total_present_marks = 0
        for row, student in enumerate(sorted_students):
            sid = student["id"]
            present = present_counts.get(sid, 0)
            total_present_marks += present
            percentage = (present / total_sessions * 100.0) if total_sessions > 0 else 0.0

            name_item = QTableWidgetItem(student["name"])
            name_item.setFont(body_font(12, QFont.Weight.Medium))
            self.attendance_table.setItem(row, 0, name_item)
            self.attendance_table.setItem(row, 1, QTableWidgetItem(student["student_id"]))
            self.attendance_table.setItem(row, 2, QTableWidgetItem(f"{present}/{total_sessions}"))

            pct_item = QTableWidgetItem(f"{percentage:.1f}%")
            pct_item.setFont(body_font(12, QFont.Weight.DemiBold))
            pct_item.setForeground(QColor(Theme.success if percentage >= 75 else Theme.warning if percentage >= 50 else Theme.danger))
            self.attendance_table.setItem(row, 3, pct_item)
            self.attendance_table.setItem(row, 4, QTableWidgetItem(student["batch_name"]))

        self.attendance_table.resizeColumnsToContents()

        student_count = len(sorted_students)
        denominator = student_count * total_sessions
        overall_percentage = (total_present_marks / denominator * 100.0) if denominator > 0 else 0.0
        self.attendance_summary_label.setText(
            f"Semester: {sem_data['name']} | Subject: {subject_name} | Students: {student_count} | Sessions: {total_sessions} | Overall Present: {overall_percentage:.1f}%"
        )

    def _generic_tab(self, title: str):
        """Create generic placeholder tab"""
        widget = CardFrame(padding=40)
        placeholder = QLabel(f"📋 {title} Report\n\nThis section will display {title.lower()} data.\nData will be loaded from backend when available.")
        placeholder.setAlignment(Qt.AlignmentFlag.AlignCenter)
        placeholder.setFont(body_font(14))
        placeholder.setStyleSheet(f"color: {Theme.text_muted};")
        placeholder.setWordWrap(True)
        widget.layout.addWidget(placeholder)
        return widget

    # =========================================================
    # TAB 2 — Viva Marks
    # =========================================================

    def _create_viva_tab(self):
        container = QWidget()
        v = QVBoxLayout(container)
        v.setContentsMargins(20, 20, 20, 20)
        v.setSpacing(14)

        top_row = QHBoxLayout()

        desc = QLabel(
            "Upload the Excel report exported from Quizizz.\n"
            "SmartLab will automatically read and display the results below."
        )
        desc.setStyleSheet(f"color: {Theme.text_muted}; font-size: 13px;")
        desc.setWordWrap(True)
        top_row.addWidget(desc, 1)

        upload_btn = QPushButton("📂  Upload Quizizz Report (.xlsx)")
        upload_btn.setStyleSheet(self._btn(Theme.primary, "#1565C0"))
        upload_btn.setFixedHeight(40)
        upload_btn.clicked.connect(self._upload_quizizz_excel)
        top_row.addWidget(upload_btn)

        v.addLayout(top_row)

        line = QFrame()
        line.setFrameShape(QFrame.Shape.HLine)
        line.setStyleSheet("color: #e0e0e0;")
        v.addWidget(line)

        self.viva_file_label = QLabel("No file loaded yet.")
        self.viva_file_label.setStyleSheet(
            f"color: {Theme.text_muted}; font-size: 12px; font-style: italic;"
        )
        v.addWidget(self.viva_file_label)

        self.viva_table = StyledTableWidget(0, 4)
        self.viva_table.setHorizontalHeaderLabels(
            ["Rank", "Student Name", "Accuracy", "Score"]
        )
        self.viva_table.setStyleSheet(f"""
            QTableWidget {{
                background: white;
                border: 1px solid {Theme.border};
                border-radius: 8px;
                gridline-color: #f3f4f6;
                font-size: 13px;
                selection-background-color: transparent;
            }}
            QHeaderView::section {{
                background: #f9fafb;
                font-weight: 700;
                font-size: 12px;
                color: #374151;
                padding: 10px 8px;
                border: none;
                border-bottom: 2px solid {Theme.border};
                text-transform: uppercase;
            }}
            QTableWidget::item {{
                padding: 10px;
            }}
            QTableWidget::item:alternate {{
                background: #fafafa;
            }}
        """)
        self.viva_table.verticalHeader().setVisible(False)
        self.viva_table.setShowGrid(False)
        self.viva_table.setAlternatingRowColors(True)
        v.addWidget(self.viva_table)

        self.viva_summary_label = QLabel("")
        self.viva_summary_label.setStyleSheet(
            f"color: {Theme.text_muted}; font-size: 12px;"
        )
        v.addWidget(self.viva_summary_label)

        return container

    def _upload_quizizz_excel(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Quizizz Excel Report",
            "",
            "Excel Files (*.xlsx *.xls);;All Files (*)",
        )
        if not file_path:
            return

        try:
            import openpyxl
        except ImportError:
            error(self, "Missing Library", "openpyxl is not installed.\n\nRun: pip install openpyxl")
            return

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active

            all_rows = [
                row for row in ws.iter_rows(values_only=True)
                if any(c is not None for c in row)
            ]

            if len(all_rows) < 2:
                warning(self, "Empty File", "The selected file appears to be empty.")
                return

            header_idx = 0
            for i, row in enumerate(all_rows):
                cells_lower = [str(c).strip().lower() for c in row if c is not None]
                if any(k in cells_lower for k in ("rank", "first name", "firstname", "name")):
                    header_idx = i
                    break

            headers = [
                str(c).strip().lower() if c is not None else ""
                for c in all_rows[header_idx]
            ]
            data_rows = all_rows[header_idx + 1:]

            col_map = self._detect_columns(headers)
            if not col_map:
                error(
                    self,
                    "Unrecognised Format",
                    "Could not find expected columns (Rank, Name, Accuracy, Score).\n\n"
                    "Please upload a standard Quizizz Excel export."
                )
                return

            self._populate_viva_table(data_rows, col_map)

            import os
            self.viva_file_label.setText(f"📄  Loaded: {os.path.basename(file_path)}")
            self.viva_file_label.setStyleSheet(
                f"color: {Theme.success}; font-size: 12px; font-weight: 600;"
            )

        except Exception as e:
            error(self, "Read Error", f"Could not read the Excel file:\n{str(e)}")

    def _detect_columns(self, headers: list) -> dict:
        keywords = {
            "rank": ["rank"],
            "name": ["first name", "firstname", "name", "student", "player"],
            "accuracy": ["accuracy", "acc", "%", "correct"],
            "score": ["score", "marks", "points", "total"],
        }
        col_map = {}
        for role, kws in keywords.items():
            for idx, h in enumerate(headers):
                if any(kw in h for kw in kws):
                    col_map[role] = idx
                    break
        return col_map if len(col_map) >= 2 else {}

    def _populate_viva_table(self, data_rows: list, col_map: dict):
        valid = [r for r in data_rows if any(c is not None for c in r)]
        self.viva_table.setRowCount(len(valid))

        for ri, row in enumerate(valid):
            def _cell(key):
                idx = col_map.get(key)
                if idx is None or idx >= len(row):
                    return ""
                v = row[idx]
                return str(v).strip() if v is not None else ""

            rank = _cell("rank")
            name = _cell("name")
            accuracy = _cell("accuracy")
            score = _cell("score")

            r_item = QTableWidgetItem(rank)
            r_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            r_item.setFont(body_font(13, QFont.Weight.Bold))
            self.viva_table.setItem(ri, 0, r_item)

            n_item = QTableWidgetItem(f"👤  {name}" if name else "—")
            n_item.setFont(body_font(13, QFont.Weight.Medium))
            self.viva_table.setItem(ri, 1, n_item)

            a_item = QTableWidgetItem(accuracy)
            a_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            a_item.setFont(body_font(13, QFont.Weight.DemiBold))
            try:
                val = float(accuracy.replace("%", "").strip())
                if val >= 75:
                    a_item.setForeground(QColor("#43A047"))
                elif val >= 50:
                    a_item.setForeground(QColor("#FF9800"))
                else:
                    a_item.setForeground(QColor("#e53935"))
            except ValueError:
                pass
            self.viva_table.setItem(ri, 2, a_item)

            s_item = QTableWidgetItem(score)
            s_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            s_item.setFont(body_font(13))
            self.viva_table.setItem(ri, 3, s_item)

        self.viva_table.resizeColumnsToContents()

        self.viva_summary_label.setText(
            f"✅  {len(valid)} student(s) loaded from Quizizz report."
        )
        self.viva_summary_label.setStyleSheet(
            f"color: {Theme.success}; font-size: 12px; font-weight: 600;"
        )

    # =========================================================
    # TAB 3 — Submissions
    # =========================================================

    def _create_submissions_tab(self):
        container = QWidget()
        layout = QVBoxLayout(container)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(16)

        top_row = QHBoxLayout()

        desc = QLabel(
            "Generate submission reports for individual students or for the full batch."
        )
        desc.setStyleSheet(f"color: {Theme.text_muted}; font-size: 13px;")
        desc.setWordWrap(True)
        top_row.addWidget(desc, 1)

        self.batch_report_btn = QPushButton("📄 Generate Full Batch Report")
        self.batch_report_btn.setStyleSheet(self._btn(Theme.primary, "#1565C0"))
        self.batch_report_btn.setFixedHeight(40)
        self.batch_report_btn.clicked.connect(self.generate_full_batch_report)
        top_row.addWidget(self.batch_report_btn)

        layout.addLayout(top_row)

        self.submissions_table = StyledTableWidget(0, 3)
        self.submissions_table.setHorizontalHeaderLabels(
            ["Student ID", "Name", "Action"]
        )
        self.submissions_table.verticalHeader().setVisible(False)
        self.submissions_table.setAlternatingRowColors(True)
        self.submissions_table.setShowGrid(False)
        self.submissions_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.submissions_table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.submissions_table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Fixed)
        layout.addWidget(self.submissions_table)

        return container

    def load_submission_students(self):
        if not hasattr(self.parent_window, 'current_batch_id') or self.parent_window.current_batch_id is None:
            self.submissions_table.setRowCount(0)
            return

        result = api_client.get_students(batch_id=self.parent_window.current_batch_id)
        if result["success"]:
            self._populate_submissions_table(result["data"])
        else:
            warning(self, "Load Failed", f"Failed to load students:\n{result['error']}")

    def _populate_submissions_table(self, students):
        self.submissions_table.clearContents()
        self.submissions_table.setRowCount(len(students))

        self.submissions_table.setColumnWidth(0, 160)
        self.submissions_table.setColumnWidth(1, 320)
        self.submissions_table.setColumnWidth(2, 260)

        for row, student in enumerate(students):
            student_id = student.get("student_id", "N/A")
            student_name = student.get("name", "Unknown")
            student_pk = student.get("id")

            id_item = QTableWidgetItem(student_id)
            id_item.setFont(body_font(12))
            self.submissions_table.setItem(row, 0, id_item)

            name_item = QTableWidgetItem(student_name)
            name_item.setFont(body_font(13, QFont.Weight.Medium))
            self.submissions_table.setItem(row, 1, name_item)

            btn = QPushButton("Generate Report")
            btn.setFixedSize(170, 34)
            btn.setCursor(Qt.CursorShape.PointingHandCursor)
            btn.setStyleSheet("""
                QPushButton {
                    background: #1abc9c;
                    color: white;
                    border: none;
                    border-radius: 6px;
                    padding: 6px 12px;
                    font-weight: 600;
                    font-size: 12px;
                    text-align: center;
                }
                QPushButton:hover {
                    background: #16a085;
                }
                QPushButton:pressed {
                    background: #12876f;
                }
            """)
            btn.clicked.connect(
                lambda checked=False, sid=student_pk, sname=student_name: self.generate_student_report(sid, sname)
            )

            wrapper = QWidget()
            wrapper.setStyleSheet("background: transparent;")
            wrapper_layout = QHBoxLayout(wrapper)
            wrapper_layout.setContentsMargins(0, 0, 0, 0)
            wrapper_layout.setSpacing(0)
            wrapper_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
            wrapper_layout.addWidget(btn)

            self.submissions_table.setCellWidget(row, 2, wrapper)
            self.submissions_table.setRowHeight(row, 56)
            self.submissions_table.setSelectionMode(self.submissions_table.SelectionMode.NoSelection)
            self.submissions_table.setFocusPolicy(Qt.FocusPolicy.NoFocus)
            self.submissions_table.setWordWrap(False)
    
    def generate_student_report(self, student_id, student_name):
        suggested_name = f"{student_name.replace(' ', '_').lower()}_submission_report.pdf"

        save_path, _ = QFileDialog.getSaveFileName(
            self,
            "Save Student Report PDF",
            suggested_name,
            "PDF Files (*.pdf)"
        )

        if not save_path:
            return

        if not save_path.lower().endswith(".pdf"):
            save_path += ".pdf"

        result = api_client.download_student_submission_report_pdf_to_path(student_id, save_path)

        if not result["success"]:
            error(
                self,
                "Download Failed",
                f"Could not generate student PDF report:\n{result['error']}"
            )
            return

        open_result = api_client.open_pdf_in_browser(result["path"])
        if not open_result["success"]:
            warning(
                self,
                "PDF Saved",
                f"PDF was saved here:\n{result['path']}\n\n"
                f"But it could not be opened in the browser:\n{open_result['error']}"
            )
        else:
            info(
                self,
                "Report Ready",
                f"Student report for {student_name} opened in browser.\n\n"
                f"Saved to:\n{result['path']}"
            )

    def generate_full_batch_report(self):
        if not hasattr(self.parent_window, 'current_batch_id') or self.parent_window.current_batch_id is None:
            info(self, "No Batch Selected", "Please select a batch first.")
            return

        batch_name = getattr(self.parent_window, "current_batch", "batch")
        suggested_name = f"{str(batch_name).replace(' ', '_').lower()}_full_submission_report.pdf"

        save_path, _ = QFileDialog.getSaveFileName(
            self,
            "Save Full Batch Report PDF",
            suggested_name,
            "PDF Files (*.pdf)"
        )

        if not save_path:
            return

        if not save_path.lower().endswith(".pdf"):
            save_path += ".pdf"

        result = api_client.download_batch_submission_report_pdf_to_path(
            self.parent_window.current_batch_id,
            save_path
        )

        if not result["success"]:
            error(
                self,
                "Download Failed",
                f"Could not generate full batch PDF report:\n{result['error']}"
            )
            return

        open_result = api_client.open_pdf_in_browser(result["path"])
        if not open_result["success"]:
            warning(
                self,
                "PDF Saved",
                f"PDF was saved here:\n{result['path']}\n\n"
                f"But it could not be opened in the browser:\n{open_result['error']}"
            )
        else:
            info(
                self,
                "Report Ready",
                f"Full batch report opened in browser.\n\n"
                f"Saved to:\n{result['path']}"
            )

    def showEvent(self, event):
        super().showEvent(event)
        self.load_reports()
